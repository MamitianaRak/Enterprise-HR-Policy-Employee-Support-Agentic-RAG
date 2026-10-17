from pathlib import Path
from typing import Iterable
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument
from app.rag.common.vectorstore import add_documents


SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return PyPDFLoader(str(path)).load()
    if suffix in {".txt", ".md"}:
        return TextLoader(str(path), encoding="utf-8").load()
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": str(path)})]
    raise ValueError(f"Unsupported file type: {suffix}")


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)
    return splitter.split_documents(list(docs))


def ingest_files(paths: Iterable[Path], domain: str, namespace: str):
    """Load, tag, chunk and index files into the domain's namespace. Returns (chunks, ids)."""
    docs = []
    for path in paths:
        for doc in load_file(path):
            doc.metadata.update(domain=domain, filename=path.name)
            docs.append(doc)
    chunks = chunk_documents(docs)
    ids = add_documents(chunks, namespace)
    return chunks, ids
